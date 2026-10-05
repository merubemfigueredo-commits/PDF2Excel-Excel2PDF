from __future__ import annotations

import hashlib
import html
import io
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import pdfplumber
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, LETTER, landscape, portrait
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    TableStyle,
)


MAX_UPLOAD_MB = 50
MAX_UPLOAD_BYTES = MAX_UPLOAD_MB * 1024 * 1024

st.set_page_config(
    page_title="Conversor PDF e Excel",
    page_icon=":material/swap_horiz:",
    layout="centered",
)


def readable_file_name(filename: str, extension: str) -> str:
    stem = Path(filename).stem.strip() or "arquivo"
    safe_stem = re.sub(r"[^\w.-]+", "_", stem, flags=re.UNICODE).strip("._")
    return f"{safe_stem or 'arquivo'}_convertido.{extension}"


def safe_sheet_title(title: str, used_titles: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", title).strip()
    cleaned = cleaned[:31] or "Dados"
    candidate = cleaned
    suffix = 2
    while candidate.casefold() in used_titles:
        marker = f"_{suffix}"
        candidate = f"{cleaned[: 31 - len(marker)]}{marker}"
        suffix += 1
    used_titles.add(candidate.casefold())
    return candidate


def rectangular_rows(rows: list[list[Any]]) -> list[list[Any]]:
    width = max((len(row) for row in rows), default=0)
    if not width:
        return []
    return [list(row) + [None] * (width - len(row)) for row in rows]


def extract_pdf_tables(pdf_bytes: bytes, use_first_row_as_header: bool) -> tuple[bytes, int, int]:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    used_titles: set[str] = set()
    sheet_count = 0
    extracted_tables = 0

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            page_tables = 0

            for table_number, table in enumerate(tables, start=1):
                rows = rectangular_rows(table or [])
                rows = [row for row in rows if any(value is not None and str(value).strip() for value in row)]
                if not rows:
                    continue

                title = f"Página {page_number}" if len(tables) == 1 else f"P{page_number}_Tabela {table_number}"
                worksheet = workbook.create_sheet(safe_sheet_title(title, used_titles))
                write_rows_to_worksheet(worksheet, rows, use_first_row_as_header)
                sheet_count += 1
                extracted_tables += 1
                page_tables += 1

            if page_tables == 0:
                text = page.extract_text() or ""
                lines = [line.strip() for line in text.splitlines() if line.strip()]
                if lines:
                    worksheet = workbook.create_sheet(
                        safe_sheet_title(f"Página {page_number} - texto", used_titles)
                    )
                    worksheet.append(["Texto extraído"])
                    for line in lines:
                        worksheet.append([line])
                    style_worksheet(worksheet, has_header=True)
                    sheet_count += 1

    if sheet_count == 0:
        raise ValueError(
            "Não encontrei texto nem tabelas neste PDF. Se o arquivo for uma digitalização "
            "ou foto, é necessário aplicar OCR antes de converter."
        )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue(), sheet_count, extracted_tables


def write_rows_to_worksheet(
    worksheet: Any, rows: list[list[Any]], use_first_row_as_header: bool
) -> None:
    for row in rows:
        worksheet.append([value if value is not None else None for value in row])
    style_worksheet(worksheet, has_header=use_first_row_as_header)


def style_worksheet(worksheet: Any, has_header: bool) -> None:
    if has_header and worksheet.max_row:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245B5A")
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_index in range(1, worksheet.max_column + 1):
        column_letter = get_column_letter(column_index)
        max_length = 0
        for cell in worksheet[column_letter]:
            if isinstance(cell.value, str):
                # PDF content is text, not an Excel formula.
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)


def excel_value_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def make_paragraph(value: Any, style: ParagraphStyle) -> Paragraph:
    text = html.escape(excel_value_to_text(value)).replace("\r\n", "\n").replace("\n", "<br/>")
    return Paragraph(text or " ", style)


def worksheet_column_widths(rows: list[list[Any]], available_width: float) -> list[float]:
    column_count = max((len(row) for row in rows), default=0)
    if column_count == 0:
        return []

    sample = rows[:250]
    weights = []
    for column_index in range(column_count):
        longest = max(
            (len(excel_value_to_text(row[column_index])) for row in sample if column_index < len(row)),
            default=8,
        )
        weights.append(min(max(longest, 8), 32) ** 0.7)
    total_weight = sum(weights) or column_count
    return [available_width * weight / total_weight for weight in weights]


def convert_workbook_to_pdf(
    workbook_bytes: bytes,
    orientation: str,
    paper_size: str,
    font_size: int,
) -> tuple[bytes, int, int]:
    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=False)
    output = io.BytesIO()

    base_size = A4 if paper_size == "A4" else LETTER
    page_size = landscape(base_size) if orientation == "Paisagem" else portrait(base_size)
    page_width, _ = page_size
    left_margin = 14 * mm
    right_margin = 14 * mm
    available_width = page_width - left_margin - right_margin

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "SheetTitle",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=14,
        leading=17,
        textColor=colors.HexColor("#183B3A"),
        spaceAfter=8,
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=font_size,
        leading=font_size + 2,
        alignment=TA_LEFT,
        spaceAfter=0,
        spaceBefore=0,
        wordWrap="CJK",
    )
    header_style = ParagraphStyle(
        "TableHeader",
        parent=cell_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    story: list[Any] = []
    total_rows = 0
    included_sheets = 0
    sheets = list(workbook.worksheets)

    for sheet_index, worksheet in enumerate(sheets):
        values = [list(row) for row in worksheet.iter_rows(values_only=True)]
        while values and not any(value is not None and str(value).strip() for value in values[-1]):
            values.pop()
        if not values:
            continue

        while values and not any(value is not None and str(value).strip() for value in values[0]):
            values.pop(0)
        if not values:
            continue

        if included_sheets:
            story.append(PageBreak())
        story.append(Paragraph(html.escape(worksheet.title), title_style))
        story.append(Spacer(1, 2 * mm))

        normalized = rectangular_rows(values)
        column_count = max((len(row) for row in normalized), default=0)
        if column_count == 0:
            continue
        table_data = [
            [
                make_paragraph(value, header_style if row_index == 0 else cell_style)
                for value in row
            ]
            for row_index, row in enumerate(normalized)
        ]
        widths = worksheet_column_widths(normalized, available_width)

        table = LongTable(
            table_data,
            colWidths=widths,
            repeatRows=1,
            hAlign="LEFT",
            splitByRow=1,
            splitInRow=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#245B5A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C8D4D2")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7F6")]),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)
        total_rows += len(normalized)
        included_sheets += 1

    workbook.close()
    if included_sheets == 0:
        raise ValueError("A planilha não contém linhas com dados para exportar.")

    def draw_page_number(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#647572"))
        canvas.drawString(left_margin, 8 * mm, "Conversor PDF e Excel")
        canvas.drawRightString(page_width - right_margin, 8 * mm, f"Página {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=left_margin,
        rightMargin=right_margin,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
        title="Planilha convertida para PDF",
    )
    document.build(
        story,
        onFirstPage=draw_page_number,
        onLaterPages=draw_page_number,
    )
    return output.getvalue(), included_sheets, total_rows


def render_cached_download(
    state_key: str,
    input_fingerprint: str,
    download_name: str,
    mime: str,
    button_label: str,
) -> None:
    result = st.session_state.get(state_key)
    if result and result["fingerprint"] == input_fingerprint:
        st.success(result["message"])
        st.download_button(
            label=button_label,
            data=result["bytes"],
            file_name=download_name,
            mime=mime,
            key=f"{state_key}_download",
            icon=":material/download:",
        )


st.title("Conversor PDF e Excel")
st.write(
    "Converta tabelas de PDF em planilhas ou exporte abas do Excel para PDF. "
    "Os arquivos são processados durante a sessão e não são armazenados pelo app."
)

pdf_tab, excel_tab = st.tabs(["PDF para Excel", "Excel para PDF"])

with pdf_tab:
    st.subheader("Converter PDF para Excel")
    st.write(
        "Cada tabela encontrada vira uma aba. Se uma página não tiver tabela detectável, "
        "o texto selecionável será colocado em uma aba própria."
    )
    pdf_file = st.file_uploader(
        "Selecione um arquivo PDF",
        type=["pdf"],
        max_upload_size=MAX_UPLOAD_MB,
        key="pdf_upload",
    )
    use_first_row_as_header = st.checkbox(
        "Usar a primeira linha de cada tabela como cabeçalho",
        value=True,
        key="pdf_header",
    )

    if pdf_file:
        pdf_bytes = pdf_file.getvalue()
        pdf_fingerprint = hashlib.sha256(
            pdf_bytes + str(use_first_row_as_header).encode("utf-8")
        ).hexdigest()
        if len(pdf_bytes) > MAX_UPLOAD_BYTES:
            st.error(f"O arquivo excede o limite de {MAX_UPLOAD_MB} MB.")
        elif b"%PDF" not in pdf_bytes[:1024]:
            st.error("O arquivo enviado não parece ser um PDF válido.")
        elif st.button(
            "Converter para .xlsx",
            type="primary",
            key="convert_pdf",
            icon=":material/table_view:",
        ):
            try:
                result_bytes, sheet_count, table_count = extract_pdf_tables(
                    pdf_bytes, use_first_row_as_header
                )
                if table_count:
                    message = (
                        f"Conversão concluída: {table_count} tabela(s) em "
                        f"{sheet_count} aba(s)."
                    )
                else:
                    message = (
                        f"Conversão concluída: texto extraído em {sheet_count} aba(s). "
                        "Nenhuma tabela estruturada foi detectada."
                    )
                st.session_state["pdf_conversion"] = {
                    "fingerprint": pdf_fingerprint,
                    "bytes": result_bytes,
                    "message": message,
                }
            except Exception as exc:
                st.session_state.pop("pdf_conversion", None)
                st.error("Não foi possível converter este PDF.")
                with st.expander("Ver detalhe do erro"):
                    st.code(str(exc))

        if b"%PDF" in pdf_bytes[:1024] and len(pdf_bytes) <= MAX_UPLOAD_BYTES:
            render_cached_download(
                "pdf_conversion",
                pdf_fingerprint,
                readable_file_name(pdf_file.name, "xlsx"),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "Baixar arquivo .xlsx",
            )
    else:
        st.info("Escolha um PDF com texto selecionável para iniciar.")

with excel_tab:
    st.subheader("Converter Excel para PDF")
    st.write(
        "Todas as abas com dados serão incluídas no mesmo PDF, cada uma começando "
        "em uma nova página. O conteúdo das células é exportado; a formatação original "
        "da planilha não é reproduzida. Fórmulas são exibidas como texto, sem cálculo."
    )
    excel_file = st.file_uploader(
        "Selecione um arquivo Excel (.xlsx)",
        type=["xlsx"],
        max_upload_size=MAX_UPLOAD_MB,
        key="excel_upload",
    )
    option_one, option_two = st.columns(2)
    with option_one:
        orientation = st.selectbox(
            "Orientação da página",
            ["Paisagem", "Retrato"],
            key="pdf_orientation",
        )
    with option_two:
        paper_size = st.selectbox(
            "Tamanho do papel",
            ["A4", "Carta"],
            key="pdf_paper_size",
        )
    font_size = st.slider("Tamanho do texto", min_value=7, max_value=12, value=9)

    if excel_file:
        excel_bytes = excel_file.getvalue()
        excel_fingerprint = hashlib.sha256(
            excel_bytes
            + f"{orientation}|{paper_size}|{font_size}".encode("utf-8")
        ).hexdigest()
        if len(excel_bytes) > MAX_UPLOAD_BYTES:
            st.error(f"O arquivo excede o limite de {MAX_UPLOAD_MB} MB.")
        elif not excel_bytes.startswith(b"PK"):
            st.error("O arquivo enviado não parece ser uma planilha .xlsx válida.")
        elif st.button(
            "Converter para PDF",
            type="primary",
            key="convert_excel",
            icon=":material/picture_as_pdf:",
        ):
            try:
                result_bytes, sheet_count, row_count = convert_workbook_to_pdf(
                    excel_bytes, orientation, paper_size, font_size
                )
                st.session_state["excel_conversion"] = {
                    "fingerprint": excel_fingerprint,
                    "bytes": result_bytes,
                    "message": (
                        f"Conversão concluída: {sheet_count} aba(s) e "
                        f"{row_count} linha(s) exportadas."
                    ),
                }
            except Exception as exc:
                st.session_state.pop("excel_conversion", None)
                st.error("Não foi possível converter esta planilha.")
                with st.expander("Ver detalhe do erro"):
                    st.code(str(exc))

        if excel_bytes.startswith(b"PK") and len(excel_bytes) <= MAX_UPLOAD_BYTES:
            render_cached_download(
                "excel_conversion",
                excel_fingerprint,
                readable_file_name(excel_file.name, "pdf"),
                "application/pdf",
                "Baixar arquivo .pdf",
            )
    else:
        st.info("Escolha uma planilha .xlsx para iniciar.")

